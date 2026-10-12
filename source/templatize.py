import docx, copy
def setr(p, i, text, clear_to=None):
    rs=p.runs
    rs[i].text=text
    if clear_to:
        for j in range(i+1, clear_to+1): rs[j].text=''
# DEMO
d=docx.Document("src/DEMO LETTER - for POWER TRAC TRACTOR.docx")
P=d.paragraphs
setr(P[3],2,'{date}',7)
p=P[5]; setr(p,8,'{village}'); setr(p,13,'{mandal}'); setr(p,18,'{district}',19)
setr(P[13],2,'{model}')
setr(P[14],3,'{serial}')
d.save("tpl/demo.docx")
# QUOTATION
d=docx.Document("src/QUOTATION.docx")
P=d.paragraphs
setr(P[3],3,'{date}',7)
setr(P[4],9,' {name}')
setr(P[5],0,'{rel}'); P[5].runs[1].text='         '[:9]; setr(P[5],8,'{care_of}')
setr(P[6],7,'{village}')
setr(P[7],7,'{mandal}')
setr(P[8],7,'{district}')
setr(P[9],2,'-  {mobile}',3)
setr(P[13],2,'{hyp}',4)
setr(P[15],2,'{words}',10)
t=d.tables[0]
def setcell(c,text):
    para=c.paragraphs[0]
    if para.runs:
        para.runs[0].text=text
        for r in para.runs[1:]: r.text=''
    else:
        # copy formatting from a reference run
        r=para.add_run(text)
        ref=t.rows[1].cells[1].paragraphs[0].runs[0]
        r.bold=ref.bold; r.font.size=ref.font.size; r.font.name=ref.font.name
for ri,pre in [(1,'1'),(2,'2')]:
    cells=t.rows[ri].cells
    for ci,tag in enumerate(['sno','item','qty','rate','amt']):
        setcell(cells[ci],'{%s%s}'%(tag,pre))
setcell(t.rows[3].cells[4],'{total}')
d.save("tpl/quotation.docx")
print("ok")
